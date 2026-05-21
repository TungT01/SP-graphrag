"""
Convert thesis_draft_v1.md to thesis_draft_v1.docx
Chinese Master's thesis formatting with python-docx
"""
import re
from pathlib import Path
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement


def set_font(run, name_cn="宋体", name_en="Times New Roman", size=12, bold=False, italic=False):
    run.font.name = name_en
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    r = run._r
    rPr = r.get_or_add_rPr()
    rFonts = OxmlElement("w:rFonts")
    rFonts.set(qn("w:eastAsia"), name_cn)
    rFonts.set(qn("w:ascii"), name_en)
    rFonts.set(qn("w:hAnsi"), name_en)
    rPr.insert(0, rFonts)


def set_para_format(para, align=WD_ALIGN_PARAGRAPH.JUSTIFY,
                    space_before=0, space_after=6, line_spacing=1.5):
    para.alignment = align
    fmt = para.paragraph_format
    fmt.space_before = Pt(space_before)
    fmt.space_after = Pt(space_after)
    fmt.line_spacing = Pt(line_spacing * 12)


def add_heading(doc, text, level):
    para = doc.add_paragraph()
    run = para.add_run(text)
    if level == 1:
        set_font(run, "黑体", "Times New Roman", 18, bold=True)
        para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        para.paragraph_format.space_before = Pt(24)
        para.paragraph_format.space_after = Pt(12)
    elif level == 2:
        set_font(run, "黑体", "Times New Roman", 15, bold=True)
        para.alignment = WD_ALIGN_PARAGRAPH.LEFT
        para.paragraph_format.space_before = Pt(18)
        para.paragraph_format.space_after = Pt(6)
    elif level == 3:
        set_font(run, "黑体", "Times New Roman", 13, bold=True)
        para.alignment = WD_ALIGN_PARAGRAPH.LEFT
        para.paragraph_format.space_before = Pt(12)
        para.paragraph_format.space_after = Pt(4)
    return para


def add_body_para(doc, text):
    """Add a normal body paragraph with mixed CN/EN inline formatting."""
    if not text.strip():
        return None
    para = doc.add_paragraph()
    set_para_format(para)
    para.paragraph_format.first_line_indent = Pt(24)
    _add_inline(para, text)
    return para


def _add_inline(para, text):
    """Parse **bold**, *italic*, $math$ markers and add runs."""
    # Strip math for display (replace $...$ with the content)
    pattern = re.compile(r'\*\*(.+?)\*\*|\*(.+?)\*|\$(.+?)\$|([^*$]+)')
    for m in pattern.finditer(text):
        bold_t, ital_t, math_t, plain_t = m.groups()
        if bold_t:
            run = para.add_run(bold_t)
            set_font(run, bold=True)
        elif ital_t:
            run = para.add_run(ital_t)
            set_font(run, italic=True)
        elif math_t:
            # Render math as italic monospace approximation
            run = para.add_run(math_t)
            set_font(run, "Courier New", "Courier New", 11, italic=True)
        elif plain_t:
            run = para.add_run(plain_t)
            set_font(run)


def add_table_from_md(doc, lines):
    """Convert a markdown table block (list of lines) to a docx table."""
    # Filter out separator rows
    rows = [l for l in lines if not re.match(r'^\s*\|[-| :]+\|\s*$', l)]
    if not rows:
        return
    parsed = []
    for row in rows:
        cells = [c.strip() for c in row.strip().strip('|').split('|')]
        parsed.append(cells)
    if not parsed:
        return
    ncols = max(len(r) for r in parsed)
    table = doc.add_table(rows=len(parsed), cols=ncols)
    table.style = 'Table Grid'
    for i, row_data in enumerate(parsed):
        for j, cell_text in enumerate(row_data):
            if j >= ncols:
                break
            cell = table.cell(i, j)
            cell.text = ''
            para = cell.paragraphs[0]
            para.alignment = WD_ALIGN_PARAGRAPH.CENTER
            # Remove markdown bold markers for table cells
            clean = re.sub(r'\*\*(.+?)\*\*', r'\1', cell_text)
            clean = re.sub(r'\$(.+?)\$', r'\1', clean)
            run = para.add_run(clean)
            is_header = (i == 0)
            set_font(run, size=10, bold=is_header)
    doc.add_paragraph()  # spacing after table


def add_code_block(doc, lines):
    """Add a code/algorithm block as a styled paragraph."""
    text = '\n'.join(lines)
    para = doc.add_paragraph()
    para.alignment = WD_ALIGN_PARAGRAPH.LEFT
    para.paragraph_format.left_indent = Cm(1)
    para.paragraph_format.space_before = Pt(6)
    para.paragraph_format.space_after = Pt(6)
    run = para.add_run(text)
    run.font.name = "Courier New"
    run.font.size = Pt(9)
    # Light gray shading
    pPr = para._p.get_or_add_pPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), 'F2F2F2')
    pPr.append(shd)


def add_formula(doc, text):
    """Add a display formula as centered paragraph."""
    para = doc.add_paragraph()
    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    para.paragraph_format.space_before = Pt(6)
    para.paragraph_format.space_after = Pt(6)
    # Strip $$ markers
    clean = text.strip().lstrip('$$').rstrip('$$').strip()
    run = para.add_run(clean)
    set_font(run, "Courier New", "Courier New", 11, italic=True)


def convert(md_path: str, docx_path: str):
    with open(md_path, encoding='utf-8') as f:
        lines = f.readlines()

    doc = Document()
    # Set margins
    for section in doc.sections:
        section.top_margin = Cm(2.54)
        section.bottom_margin = Cm(2.54)
        section.left_margin = Cm(3.0)
        section.right_margin = Cm(2.5)

    i = 0
    table_buffer = []
    code_buffer = []
    in_code = False
    formula_buffer = []
    in_formula = False

    while i < len(lines):
        line = lines[i].rstrip('\n')

        # Code block
        if line.strip().startswith('```'):
            if in_code:
                add_code_block(doc, code_buffer)
                code_buffer = []
                in_code = False
            else:
                in_code = True
            i += 1
            continue
        if in_code:
            code_buffer.append(line)
            i += 1
            continue

        # Display formula $$...$$
        if line.strip().startswith('$$') and not line.strip()[2:].strip():
            if in_formula:
                add_formula(doc, '\n'.join(formula_buffer))
                formula_buffer = []
                in_formula = False
            else:
                in_formula = True
            i += 1
            continue
        if in_formula:
            formula_buffer.append(line)
            i += 1
            continue

        # Single-line $$formula$$
        m_formula = re.match(r'^\s*\$\$(.+)\$\$\s*$', line)
        if m_formula:
            add_formula(doc, m_formula.group(1))
            i += 1
            continue

        # Markdown table — collect all table lines
        if line.strip().startswith('|'):
            table_buffer.append(line)
            i += 1
            # peek for more table lines
            while i < len(lines) and lines[i].strip().startswith('|'):
                table_buffer.append(lines[i].rstrip('\n'))
                i += 1
            add_table_from_md(doc, table_buffer)
            table_buffer = []
            continue

        # Headings
        if re.match(r'^# ', line):
            add_heading(doc, line[2:].strip(), 1)
            i += 1
            continue
        if re.match(r'^## ', line):
            add_heading(doc, line[3:].strip(), 2)
            i += 1
            continue
        if re.match(r'^### ', line):
            add_heading(doc, line[4:].strip(), 3)
            i += 1
            continue

        # Horizontal rule
        if re.match(r'^---+\s*$', line):
            para = doc.add_paragraph()
            para.paragraph_format.border_bottom = True
            i += 1
            continue

        # Blockquote (> ...)
        if line.strip().startswith('>'):
            content = re.sub(r'^>\s*', '', line.strip())
            if content:
                para = doc.add_paragraph()
                para.alignment = WD_ALIGN_PARAGRAPH.LEFT
                para.paragraph_format.left_indent = Cm(1)
                para.paragraph_format.space_before = Pt(4)
                para.paragraph_format.space_after = Pt(4)
                run = para.add_run(content)
                set_font(run, size=10.5, italic=True)
            i += 1
            continue

        # Empty line
        if not line.strip():
            i += 1
            continue

        # Reference list items [N] ...
        if re.match(r'^\[\d+\]', line.strip()):
            para = doc.add_paragraph()
            para.alignment = WD_ALIGN_PARAGRAPH.LEFT
            para.paragraph_format.left_indent = Cm(0.8)
            para.paragraph_format.first_line_indent = Pt(-18)
            para.paragraph_format.space_after = Pt(4)
            run = para.add_run(line.strip())
            set_font(run, size=10.5)
            i += 1
            continue

        # Bold-only lines like **贡献一**：...
        # Just treat as body paragraph
        add_body_para(doc, line.strip())
        i += 1

    doc.save(docx_path)
    print(f"Saved: {docx_path}")


if __name__ == '__main__':
    base = Path("/Users/ttung/Desktop/个人学习/SP-GraphRAG/reports")
    convert(str(base / "thesis_draft_v1.md"), str(base / "thesis_draft_v1.docx"))
